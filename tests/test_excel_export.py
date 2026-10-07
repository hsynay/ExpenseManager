"""Excel downloads on /payments, /checks, /expenses and /debts.

What is checked for every page: the file downloads, it holds the rows the
filters and the search select (not just the first page), the money columns are
numbers and the dates are dates. The page and the file must also agree, so
those tests read the rows the page handed to its template and compare.
"""
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

import pytest
from openpyxl import load_workbook

import app as app_module
from excel_export import ExcelBook, TEXT, MONEY, DATE, XLSX_MIME
from helpers import PAGE_SIZE

D = Decimal
JUNE = date(2026, 6, 15)


def download(client, url):
    response = client.get(url)
    assert response.status_code == 200, (url, response.status_code)
    assert response.mimetype == XLSX_MIME
    return response


def open_book(response):
    return load_workbook(BytesIO(response.data))


def data_rows(book, sheet):
    """The rows under the header row, as lists of plain values."""
    return [list(row) for row in
            book[sheet].iter_rows(min_row=2, values_only=True)]


def owner_with_flat(factory, project_id, first, flat_no=1, price=D('100000')):
    customer_id = factory.customer(first=first)
    flat_id = factory.flat(project_id, owner_id=customer_id, flat_no=flat_no,
                           total_price=price)
    return customer_id, flat_id


# --- every export needs a login ---------------------------------------------

@pytest.mark.parametrize('url', [
    '/payments/export', '/checks/export', '/expenses/export?project_id=1',
    '/debts/export'])
def test_exports_need_a_login(url):
    with app_module.app.test_client() as anonymous:
        response = anonymous.get(url)

    assert response.status_code == 302
    assert '/login' in response.headers['Location']


# --- the builder itself -----------------------------------------------------

def test_builder_formats_header_money_and_dates():
    book = ExcelBook()
    book.add_sheet('Deneme', [('Ad', TEXT), ('Tutar', MONEY), ('Tarih', DATE)],
                   [['Ali', D('1234.5'), date(2026, 3, 9)], ['Veli', None, None]])
    sheet = open_book(_as_response(book))['Deneme']

    assert all(sheet.cell(1, col).font.bold for col in (1, 2, 3))
    assert sheet['B2'].value == 1234.5
    assert sheet['B2'].number_format == '#,##0.00'
    assert sheet['C2'].value == datetime(2026, 3, 9)
    assert sheet['C2'].number_format == 'DD.MM.YYYY'
    # An empty value stays an empty cell.
    assert sheet['B3'].value is None and sheet['C3'].value is None


def test_builder_never_turns_text_into_a_formula():
    book = ExcelBook()
    book.add_sheet('Deneme', [('Ad', TEXT)], [['=1+1'], ['@SUM(A1)'], ['-5']])
    sheet = open_book(_as_response(book))['Deneme']

    for row in (2, 3, 4):
        assert sheet.cell(row, 1).data_type == 's'
    assert sheet['A2'].value == '=1+1'


def test_builder_refuses_a_row_of_the_wrong_length():
    book = ExcelBook()
    with pytest.raises(ValueError):
        book.add_sheet('Deneme', [('A', TEXT), ('B', TEXT)], [['only one']])


def _as_response(book):
    with app_module.app.test_request_context():
        response = book.response('deneme')
    # A file response streams by default; the test wants the bytes.
    response.direct_passthrough = False
    return response


# --- /payments --------------------------------------------------------------

def test_payments_export_is_named_formatted_and_complete(client, factory,
                                                         render_context):
    project_id = factory.project('xlsx gelir')
    _, flat_id = owner_with_flat(factory, project_id, 'XlsxAli')
    count = PAGE_SIZE + 2
    for n in range(count):
        factory.payment(flat_id, D('1000.50'),
                        payment_date=JUNE + timedelta(days=n % 5),
                        description='xlsx %d' % n)
    # Another project, to show the filter keeps it out.
    other_project = factory.project('xlsx gelir baska')
    _, other_flat = owner_with_flat(factory, other_project, 'XlsxBaska')
    factory.payment(other_flat, D('1'), description='baska')
    factory.commit()
    args = '?project=' + quote('PYTEST xlsx gelir')

    client.get('/payments' + args)
    assert len(render_context.latest('payments.html')['payments']) == PAGE_SIZE

    response = download(client, '/payments/export' + args)
    assert response.headers['Content-Disposition'] == (
        'attachment; filename=payments_%s.xlsx' % date.today().isoformat())
    book = open_book(response)
    rows = data_rows(book, 'Gelirler')

    # The page shows 50 rows, the file holds all of them.
    assert len(rows) == count
    sheet = book['Gelirler']
    assert [c.value for c in sheet[1]] == [
        'Proje', 'Müşteri', 'Daire', 'Ödeme Yöntemi', 'Açıklama',
        'Ücret (₺)', 'Tarih']
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet['F2'].number_format == '#,##0.00'
    assert sheet['G2'].number_format == 'DD.MM.YYYY'
    expected = ['PYTEST xlsx gelir', 'PYTEST XlsxAli Musteri',
                'Blok: A, Kat: 1, No: 1', 'Nakit', 'PYTEST xlsx 7', 1000.5,
                datetime(2026, 6, 17)]
    assert expected in rows


def test_payments_export_follows_date_range_and_search(client, factory):
    project_id = factory.project('xlsx gelir filtre')
    _, flat_a = owner_with_flat(factory, project_id, 'XlsxZehra')
    _, flat_b = owner_with_flat(factory, project_id, 'XlsxMelek', flat_no=2)
    for n in range(10):
        factory.payment(flat_a, D('10'), payment_date=JUNE + timedelta(days=n))
    for n in range(4):
        factory.payment(flat_b, D('20'), payment_date=JUNE + timedelta(days=n))
    factory.commit()
    base = '/payments/export?project=' + quote('PYTEST xlsx gelir filtre')

    assert len(data_rows(open_book(download(client, base)), 'Gelirler')) == 14

    after = base + '&start_date=' + (JUNE + timedelta(days=6)).isoformat()
    assert len(data_rows(open_book(download(client, after)), 'Gelirler')) == 4

    searched = base + '&search=XlsxMelek'
    rows = data_rows(open_book(download(client, searched)), 'Gelirler')
    assert len(rows) == 4
    assert all('XlsxMelek' in row[1] for row in rows)


def test_payments_export_matches_the_page_when_sorted(client, factory,
                                                      render_context):
    project_id = factory.project('xlsx gelir sira')
    _, flat_id = owner_with_flat(factory, project_id, 'XlsxSira')
    for amount in ('30', '10', '20'):
        factory.payment(flat_id, D(amount))
    factory.commit()
    args = ('?project=' + quote('PYTEST xlsx gelir sira')
            + '&sort_by=tutar&order=asc')

    client.get('/payments' + args)
    on_page = [p[7] for p in render_context.latest('payments.html')['payments']]
    in_file = [row[5] for row in
               data_rows(open_book(download(client, '/payments/export' + args)),
                         'Gelirler')]

    assert in_file == [float(a) for a in on_page] == [10.0, 20.0, 30.0]


# --- /checks ----------------------------------------------------------------

def test_checks_export_filters_each_list_and_labels_the_status(client, factory,
                                                               render_context):
    customer_id = factory.customer(first='XlsxCekci')
    factory.check(customer_id, D('500'), status='tahsil_edildi',
                  due_date=date(2026, 7, 1))
    factory.check(customer_id, D('700.25'), status='portfoyde',
                  due_date=date(2026, 8, 1))
    factory.check(customer_id, D('900'), status='karsiliksiz',
                  due_date=date(2026, 9, 1))
    supplier_id = factory.supplier(name='XlsxTedarikci')
    factory.outgoing_check(supplier_id, D('300'), status='verildi',
                           due_date=date(2026, 7, 5))
    factory.outgoing_check(supplier_id, D('400'), status='odendi',
                           due_date=date(2026, 7, 6))
    factory.commit()
    parties = 'in_customer=XlsxCekci&out_supplier=XlsxTedarikci'

    everything = open_book(download(client, '/checks/export?' + parties))
    assert len(data_rows(everything, 'Alınan Çekler')) == 3
    assert len(data_rows(everything, 'Verilen Çekler')) == 2

    narrowed = '/checks/export?%s&in_status=tahsil_edildi&out_status=odendi' % parties
    book = open_book(download(client, narrowed))
    incoming = data_rows(book, 'Alınan Çekler')
    outgoing = data_rows(book, 'Verilen Çekler')
    assert [r[4:] for r in incoming] == [[500, 'Tahsil Edildi']]
    assert [r[4:] for r in outgoing] == [[400, 'Ödendi']]
    assert incoming[0][0] == datetime(2026, 7, 1)
    assert incoming[0][1] == 'PYTEST XlsxCekci Musteri'

    # The page, given the same query string, shows the same rows.
    client.get('/checks?%s&in_status=tahsil_edildi&out_status=odendi' % parties)
    page = render_context.latest('checks.html')
    assert len(page['incoming_checks']) == len(incoming)
    assert len(page['outgoing_checks']) == len(outgoing)

    book = open_book(download(client, '/checks/export?%s&in_status=karsiliksiz'
                              % parties))
    assert [r[5] for r in data_rows(book, 'Alınan Çekler')] == ['Karşılıksız']
    assert book['Alınan Çekler']['E2'].number_format == '#,##0.00'
    assert book['Alınan Çekler']['A2'].number_format == 'DD.MM.YYYY'


def test_checks_export_has_every_row_not_one_page(client, factory,
                                                  render_context):
    customer_id = factory.customer(first='XlsxCokCek')
    for n in range(PAGE_SIZE + 2):
        factory.check(customer_id, D('10'), due_date=JUNE + timedelta(days=n))
    factory.commit()
    args = '?in_customer=XlsxCokCek'

    client.get('/checks' + args)
    assert len(render_context.latest('checks.html')['incoming_checks']) == PAGE_SIZE

    book = open_book(download(client, '/checks/export' + args))
    assert len(data_rows(book, 'Alınan Çekler')) == PAGE_SIZE + 2


def test_checks_export_with_a_broken_date_goes_back_to_the_page(client):
    response = client.get('/checks/export?in_due_from=not-a-date')

    assert response.status_code == 302
    assert response.headers['Location'].endswith('/checks')


# --- /expenses --------------------------------------------------------------

def expense_project(factory):
    """A project with two suppliers, two expenses and two petty cash rows."""
    project_id = factory.project('xlsx gider')
    supplier_a = factory.supplier(project_id, name='XlsxA')
    supplier_b = factory.supplier(project_id, name='XlsxB')
    beton = factory.expense(project_id, supplier_a, amount=D('10000'),
                            title='xlsx beton')
    factory.expense_installment(beton, JUNE, D('5000'))
    factory.expense_installment(beton, JUNE + timedelta(days=30), D('5000'))
    factory.supplier_payment(beton, supplier_a, D('5000'), payment_date=JUNE)
    factory.expense(project_id, supplier_b, amount=D('2000'), title='xlsx demir')
    factory.petty_cash(project_id, D('50'), title='xlsx fis')
    factory.petty_cash(project_id, D('75.5'), title='xlsx yemek',
                       expense_date=JUNE + timedelta(days=1))
    factory.commit()
    return project_id, supplier_a, supplier_b


def test_expenses_export_has_a_sheet_per_table(client, factory, render_context):
    project_id, supplier_a, _ = expense_project(factory)

    client.get('/expenses?project_id=%d' % project_id)
    page = render_context.latest('expenses.html')

    book = open_book(download(client, '/expenses/export?project_id=%d' % project_id))
    assert book.sheetnames == ['Büyük Giderler', 'Gider Taksitleri',
                               'Gider Ödemeleri', 'Küçük Giderler']

    big = data_rows(book, 'Büyük Giderler')
    assert len(big) == len(page['expenses_data']) == 2
    beton = next(r for r in big if r[0] == 'PYTEST xlsx beton')
    assert beton[2] == 10000 and isinstance(beton[3], (int, float))

    assert len(data_rows(book, 'Gider Taksitleri')) == 2
    payments = data_rows(book, 'Gider Ödemeleri')
    assert len(payments) == 1
    assert payments[0][0] == 'PYTEST xlsx beton'
    assert payments[0][2] == datetime(2026, 6, 15)
    assert payments[0][4] == 'Nakit' and payments[0][9] == 5000

    petty = data_rows(book, 'Küçük Giderler')
    assert len(petty) == len(page['petty_cash_items']) == 2
    # Newest first, like the page.
    assert petty[0] == [datetime(2026, 6, 16), 'PYTEST xlsx yemek', 75.5]

    assert book['Büyük Giderler']['C2'].number_format == '#,##0.00'
    assert book['Gider Taksitleri']['C2'].number_format == 'DD.MM.YYYY'


def test_expenses_export_follows_the_filters(client, factory):
    project_id, supplier_a, _ = expense_project(factory)
    base = '/expenses/export?project_id=%d' % project_id

    by_supplier = open_book(download(client, base + '&supplier_id=%d' % supplier_a))
    assert [r[0] for r in data_rows(by_supplier, 'Büyük Giderler')] == [
        'PYTEST xlsx beton']
    assert len(data_rows(by_supplier, 'Gider Taksitleri')) == 2

    by_title = open_book(download(client, base + '&title=demir'))
    assert [r[0] for r in data_rows(by_title, 'Büyük Giderler')] == [
        'PYTEST xlsx demir']
    assert data_rows(by_title, 'Küçük Giderler') == []

    only_petty = open_book(download(client, base + '&expense_type=petty'))
    assert only_petty.sheetnames == ['Küçük Giderler']

    only_large = open_book(download(client, base + '&expense_type=large'))
    assert 'Küçük Giderler' not in only_large.sheetnames


def test_expenses_export_has_every_row_not_one_page(client, factory,
                                                    render_context):
    project_id = factory.project('xlsx gider cok')
    for n in range(PAGE_SIZE + 2):
        factory.expense(project_id, amount=D('10'), title='xlsx toplu %d' % n)
    factory.commit()

    client.get('/expenses?project_id=%d' % project_id)
    assert len(render_context.latest('expenses.html')['expenses_data']) == PAGE_SIZE

    book = open_book(download(client, '/expenses/export?project_id=%d' % project_id))
    assert len(data_rows(book, 'Büyük Giderler')) == PAGE_SIZE + 2


def test_expenses_export_without_a_project_goes_back_to_the_page(client):
    response = client.get('/expenses/export')

    assert response.status_code == 302
    assert response.headers['Location'].endswith('/expenses')


# --- /debts -----------------------------------------------------------------

def test_debts_export_matches_the_page(client, factory, render_context):
    project_id = factory.project('xlsx borc')
    customer_a, flat_a = owner_with_flat(factory, project_id, 'XlsxBir',
                                         price=D('100000'))
    _, flat_b = owner_with_flat(factory, project_id, 'XlsxIki', flat_no=2,
                                price=D('50000'))
    for month in range(3):
        factory.installment(flat_a, JUNE + timedelta(days=30 * month), D('1000'))
    factory.installment(flat_b, JUNE, D('2000'))
    factory.payment(flat_a, D('1000'))
    # A check still in the portfolio is listed but does not count as paid.
    pending = factory.check(customer_a, D('700'), status='portfoyde')
    factory.payment(flat_a, D('700'), payment_method='çek', check_id=pending)
    factory.commit()
    args = '?project_id=%d' % project_id

    client.get('/debts' + args)
    page_flats = [f for p in render_context.latest('debts.html')['projects_data']
                  for f in p['flats']]
    book = open_book(download(client, '/debts/export' + args))

    flats = data_rows(book, 'Daireler')
    assert len(flats) == len(page_flats) == 2
    assert len(data_rows(book, 'Taksitler')) == 4
    payments = data_rows(book, 'Ödemeler')
    assert len(payments) == 2

    assert data_rows(book, 'Projeler') == [
        ['PYTEST xlsx borc', 'Normal', 150000, 1000, 149000]]
    first = next(r for r in flats if 'XlsxBir' in r[1])
    assert first[3:] == [100000, 1000, 99000]
    check_row = next(r for r in payments if r[5] == 'Çek')
    assert check_row[6] == 'Portföyde' and check_row[10] == 700

    assert book['Daireler']['D2'].number_format == '#,##0.00'
    assert book['Taksitler']['D2'].number_format == 'DD.MM.YYYY'
    assert [c.value for c in book['Daireler'][1]][:3] == [
        'Proje', 'Müşteri', 'Daire']


def test_debts_export_follows_search_and_flat_filters(client, factory):
    project_id = factory.project('xlsx borc arama')
    _, flat_a = owner_with_flat(factory, project_id, 'XlsxAranan')
    _, flat_b = owner_with_flat(factory, project_id, 'XlsxDigeri', flat_no=2)
    factory.installment(flat_a, JUNE, D('100'))
    factory.installment(flat_a, JUNE + timedelta(days=30), D('100'))
    factory.installment(flat_b, JUNE, D('100'))
    factory.commit()
    base = '/debts/export?project_id=%d' % project_id

    searched = open_book(download(client, base + '&search=XlsxAranan'))
    assert len(data_rows(searched, 'Daireler')) == 1
    assert len(data_rows(searched, 'Taksitler')) == 2

    one_flat = open_book(download(client, base + '&flat_id=%d' % flat_b))
    assert [r[1] for r in data_rows(one_flat, 'Daireler')] == [
        'PYTEST XlsxDigeri Musteri']
    assert len(data_rows(one_flat, 'Taksitler')) == 1

    nobody = open_book(download(client, base + '&search=XlsxYokBoyleBiri'))
    assert data_rows(nobody, 'Daireler') == []
    assert data_rows(nobody, 'Projeler') == []


def test_debts_export_of_a_cooperative_shows_only_what_was_paid(client, factory):
    project_id = factory.project('xlsx koop', project_type='cooperative')
    _, flat_id = owner_with_flat(factory, project_id, 'XlsxUye', price=None)
    factory.payment(flat_id, D('200'))
    factory.commit()

    book = open_book(download(client, '/debts/export?project_id=%d' % project_id))

    assert data_rows(book, 'Projeler') == [
        ['PYTEST xlsx koop', 'Kooperatif', None, None, None]]
    assert data_rows(book, 'Daireler') == [
        ['PYTEST xlsx koop', 'PYTEST XlsxUye Musteri',
         'Blok: A, Kat: 1, No: 1', None, 200, None]]
    assert data_rows(book, 'Taksitler') == []
    assert len(data_rows(book, 'Ödemeler')) == 1
