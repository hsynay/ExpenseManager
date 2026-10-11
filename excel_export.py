"""Excel downloads for the list pages.

Every export looks the same: a bold header row, money columns stored as real
numbers with a thousands format (so Excel can still add them up) and dates
shown as DD.MM.YYYY.

Usage:
    book = ExcelBook()
    book.add_sheet('Odemeler', [('Proje', TEXT), ('Tutar', MONEY)], rows)
    return book.response('Odemeler')
"""
import re
import unicodedata
from io import BytesIO
from urllib.parse import quote

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from helpers import today_istanbul

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MONEY_FORMAT = '#,##0.00'
DATE_FORMAT = 'DD.MM.YYYY'

# Column kinds
TEXT, MONEY, DATE = 'text', 'money', 'date'

MAX_COLUMN_WIDTH = 60
HEADER_FONT = Font(bold=True)


class ExcelBook:
    """One .xlsx file made of one or more sheets."""

    def __init__(self):
        self._workbook = Workbook()
        # A new workbook starts with one empty sheet. The first add_sheet
        # call renames it instead of leaving it behind.
        self._unused_first_sheet = True

    def add_sheet(self, title, columns, rows):
        """Write one sheet.

        columns is a list of (header, kind) pairs, rows is a list of
        sequences in the same order. A row with the wrong length raises an
        error, so a column that drifts out of step cannot go unnoticed.
        """
        if self._unused_first_sheet:
            sheet = self._workbook.active
            sheet.title = title
            self._unused_first_sheet = False
        else:
            sheet = self._workbook.create_sheet(title)

        widths = [len(header) for header, _ in columns]
        for col, (header, _) in enumerate(columns, start=1):
            sheet.cell(row=1, column=col, value=header).font = HEADER_FONT

        for row_number, row in enumerate(rows, start=2):
            if len(row) != len(columns):
                raise ValueError('Excel row has %d values for %d columns'
                                 % (len(row), len(columns)))
            for col, ((_, kind), value) in enumerate(zip(columns, row), start=1):
                cell = sheet.cell(row=row_number, column=col)
                shown = self._fill(cell, kind, value)
                widths[col - 1] = max(widths[col - 1], len(shown))

        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = \
                min(width + 2, MAX_COLUMN_WIDTH)

    @staticmethod
    def _fill(cell, kind, value):
        """Put one value in a cell. Returns the text used to size the column."""
        if value is None:
            return ''
        if kind == MONEY:
            cell.value = value
            cell.number_format = MONEY_FORMAT
            return '{:,.2f}'.format(value)
        if kind == DATE:
            cell.value = value
            cell.number_format = DATE_FORMAT
            return 'DD.MM.YYYY'
        text = str(value)
        cell.value = text
        # openpyxl turns any text that starts with "=" into a formula. Names
        # and descriptions are typed by users, so force plain text.
        cell.data_type = 's'
        return text

    def response(self, page_label, project_name=None, sub_sheet=None):
        """The finished file as a download.

        The name is Project-Name_PageLabel-SubSheet_YYYY-MM-DD.xlsx. The
        project and the sub sheet are left out when they do not apply, and the
        date is the Istanbul date.
        """
        buffer = BytesIO()
        self._workbook.save(buffer)
        buffer.seek(0)
        filename = download_filename(page_label, project_name, sub_sheet)
        response = send_file(buffer, mimetype=XLSX_MIME, as_attachment=True,
                             download_name=filename)
        # Send both forms: a plain ASCII name for old clients and the real
        # name (with Turkish letters) in the RFC 5987 form that browsers use.
        response.headers['Content-Disposition'] = (
            'attachment; filename="%s"; filename*=UTF-8\'\'%s'
            % (ascii_fallback(filename), quote(filename, safe='')))
        # Financial data: no browser or proxy should keep a copy.
        response.headers['Cache-Control'] = 'no-store'
        return response


# --- file names -------------------------------------------------------------

# Characters Windows and most file systems do not allow in a name, plus the
# control characters.
_FORBIDDEN_IN_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_TURKISH_TO_ASCII = str.maketrans('çğıöşüÇĞİÖŞÜ', 'cgiosuCGIOSU')

MAX_PROJECT_PART = 60
MAX_LABEL_PART = 40


def clean_name_part(text, max_length):
    """Make free text safe inside a file name: no forbidden characters,
    spaces become dashes, and the length is cut."""
    # Whitespace first: a tab or a line break is a control character, and
    # deleting it would glue two words together.
    text = re.sub(r'\s+', '-', (text or '').strip())
    text = _FORBIDDEN_IN_NAME.sub('', text)
    text = re.sub(r'-{2,}', '-', text)
    return text[:max_length].strip('-.')


def download_filename(page_label, project_name=None, sub_sheet=None):
    """Project-Name_PageLabel-SubSheet_YYYY-MM-DD.xlsx (parts may be absent)."""
    label = clean_name_part(page_label, MAX_LABEL_PART)
    if sub_sheet:
        label += '-' + clean_name_part(sub_sheet, MAX_LABEL_PART)
    parts = [clean_name_part(project_name, MAX_PROJECT_PART), label,
             today_istanbul().isoformat()]
    return '_'.join(p for p in parts if p) + '.xlsx'


def ascii_fallback(filename):
    """The same name with Turkish letters made plain, for the filename= form."""
    plain = filename.translate(_TURKISH_TO_ASCII)
    plain = unicodedata.normalize('NFKD', plain).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Za-z0-9._-]', '', plain) or 'download.xlsx'
